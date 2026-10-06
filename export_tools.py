"""
Export conversation to various formats.
Requires:
  pip install reportlab python-docx
"""

import os
import datetime


def _clean(text):
    if text.startswith("AI: "):
        return text[4:]
    return text


def export_txt(bubbles, path, title="Conversation"):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(f"{title}\n")
        f.write(f"Exported: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write("=" * 60 + "\n\n")
        for kind, text in bubbles:
            t = _clean(text)
            if kind == "user":
                f.write(f"You: {t}\n\n")
            elif kind == "agent":
                f.write(f"Agent: {t}\n\n")
            elif kind == "scheduled":
                f.write(f"[Scheduled] {t}\n\n")
            else:
                f.write(f"— {t}\n\n")
    return path


def export_md(bubbles, path, title="Conversation"):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(f"# {title}\n\n")
        f.write(f"_Exported: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}_\n\n")
        f.write("---\n\n")
        for kind, text in bubbles:
            t = _clean(text)
            if kind == "user":
                f.write(f"**You:** {t}\n\n")
            elif kind == "agent":
                f.write(f"**Agent:** {t}\n\n")
            elif kind == "scheduled":
                f.write(f"> ⏰ {t}\n\n")
            else:
                f.write(f"_{t}_\n\n")
    return path


def export_pdf(bubbles, path, title="Conversation"):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except ImportError:
        raise RuntimeError("reportlab not installed. Run: pip install reportlab")

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    user_style = ParagraphStyle(name="U", parent=styles["BodyText"],
                                fontSize=11, leading=15,
                                textColor=colors.HexColor("#1a73e8"),
                                fontName="Helvetica-Bold")
    agent_style = ParagraphStyle(name="A", parent=styles["BodyText"],
                                 fontSize=11, leading=15,
                                 textColor=colors.HexColor("#111111"))
    sys_style = ParagraphStyle(name="S", parent=styles["BodyText"],
                               fontSize=9, leading=12,
                               textColor=colors.HexColor("#888888"),
                               fontName="Helvetica-Oblique")

    def esc(t):
        return (t.replace("&", "&amp;")
                 .replace("<", "&lt;")
                 .replace(">", "&gt;")
                 .replace("\n", "<br/>"))

    story = [Paragraph(title, styles["Title"]),
             Paragraph(f"Exported: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", sys_style),
             Spacer(1, 0.5*cm)]

    for kind, text in bubbles:
        t = _clean(text)
        if kind == "user":
            story.append(Paragraph(f"You: {esc(t)}", user_style))
        elif kind == "agent":
            story.append(Paragraph(f"Agent: {esc(t)}", agent_style))
        elif kind == "scheduled":
            story.append(Paragraph(f"⏰ {esc(t)}", sys_style))
        else:
            story.append(Paragraph(esc(t), sys_style))
        story.append(Spacer(1, 0.3*cm))

    doc.build(story)
    return path


def export_docx(bubbles, path, title="Conversation"):
    try:
        from docx import Document
        from docx.shared import RGBColor
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Exported: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    for kind, text in bubbles:
        t = _clean(text)
        if kind == "user":
            p = doc.add_paragraph()
            r = p.add_run("You: ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x1a, 0x73, 0xe8)
            p.add_run(t)
        elif kind == "agent":
            p = doc.add_paragraph()
            r = p.add_run("Agent: ")
            r.bold = True
            p.add_run(t)
        elif kind == "scheduled":
            p = doc.add_paragraph()
            r = p.add_run("⏰ " + t)
            r.italic = True
        else:
            p = doc.add_paragraph()
            r = p.add_run(t)
            r.italic = True

    doc.save(path)
    return path