"""
Read uploaded files and turn them into text the agent can process.

Text:   .txt .md .py .js .ts .json .csv .log .html .xml
        .yaml .yml .ini .cfg .sh .bat .ps1 .sql
Rich:   .pdf   (pypdf)
        .docx  (python-docx)
        .xlsx  (openpyxl)
Images: .png .jpg .jpeg .bmp .gif .webp (OCR via screen.py or pytesseract)

Requires: pip install pypdf python-docx openpyxl pytesseract pillow
(If you already have screen.py working for OCR, that path is tried first.)
"""

import os

MAX_CHARS_PER_FILE = 50_000   # guard against context blowout

TEXT_EXTS = {
    ".txt", ".md", ".py", ".js", ".ts", ".json", ".csv", ".log",
    ".html", ".htm", ".xml", ".yaml", ".yml", ".ini", ".cfg",
    ".sh", ".bat", ".ps1", ".sql",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tiff"}


def _truncate(text):
    if len(text) <= MAX_CHARS_PER_FILE:
        return text
    extra = len(text) - MAX_CHARS_PER_FILE
    return text[:MAX_CHARS_PER_FILE] + f"\n\n… [truncated, {extra:,} more chars]"


def _read_text(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[PDF reader not installed. Run: pip install pypdf]"
    try:
        reader = PdfReader(path)
        pages = []
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                pages.append(f"--- page {i} ---\n{t}")
        return "\n\n".join(pages) if pages else "[PDF had no extractable text]"
    except Exception as e:
        return f"[PDF read failed: {e}]"


def _read_docx(path):
    try:
        from docx import Document
    except ImportError:
        return "[DOCX reader not installed. Run: pip install python-docx]"
    try:
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[DOCX read failed: {e}]"


def _read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        return "[XLSX reader not installed. Run: pip install openpyxl]"
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        chunks = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    rows.append("\t".join("" if c is None else str(c) for c in row))
            if rows:
                chunks.append(f"--- sheet: {ws.title} ---\n" + "\n".join(rows))
        return "\n\n".join(chunks) if chunks else "[XLSX was empty]"
    except Exception as e:
        return f"[XLSX read failed: {e}]"


def _read_image(path):
    """OCR an image. Tries screen.py's OCR path first, then pytesseract."""
    # Try the project's existing OCR helper first
    try:
        import screen
        for fn_name in ("ocr_image", "read_image", "extract_text", "ocr_file",
                        "describe_screen", "read_screen"):
            fn = getattr(screen, fn_name, None)
            if callable(fn):
                try:
                    out = fn(path)
                    if isinstance(out, str) and out.strip():
                        return out
                except TypeError:
                    # function probably takes no args (screen-wide OCR); skip
                    continue
                except Exception:
                    continue
    except ImportError:
        pass

    # Fallback: direct pytesseract
    try:
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(path))
    except ImportError:
        return "[image OCR unavailable: install pytesseract + pillow]"
    except Exception as e:
        return f"[image OCR failed: {e}]"


def read_file(path):
    """Return a formatted string containing the file's extracted content."""
    if not os.path.isfile(path):
        return f"[not found: {path}]"
    name = os.path.basename(path)
    size = os.path.getsize(path)
    ext = os.path.splitext(path)[1].lower()
    header = f"--- {name} ({size:,} bytes) ---"

    try:
        if ext in TEXT_EXTS:
            body = _read_text(path)
        elif ext == ".pdf":
            body = _read_pdf(path)
        elif ext == ".docx":
            body = _read_docx(path)
        elif ext == ".xlsx":
            body = _read_xlsx(path)
        elif ext in IMAGE_EXTS:
            body = _read_image(path)
        else:
            body = f"[unsupported file type: {ext}]"
    except Exception as e:
        body = f"[read failed: {e}]"

    return f"{header}\n{_truncate(body)}\n--- end {name} ---"


def _describe_image_via_agent(path, question=None):
    """Route an image through agent.describe_image() (DeepSeek vision).

    Imported lazily — agent.py imports this module, so a top-level import
    would be circular.  Returns an 'ERROR: ...' string on any failure.
    """
    try:
        import agent
    except Exception as e:
        return f"ERROR: vision unavailable: {e}"
    q = question or ("Describe this image in detail. Include any visible text, "
                     "objects, colours, and layout.")
    try:
        return agent.describe_image(f"{path}|{q}")
    except Exception as e:
        return f"ERROR: vision failed: {e}"


def build_attachment_block(paths):
    """Concatenate all attached files into a single text block for the LLM.

    Images go through vision (describe_image) first, falling back to OCR.
    Failures are reported per-file with a specific reason instead of a vague
    generic message.
    """
    if not paths:
        return ""

    parts = []
    for i, p in enumerate(paths, 1):
        name = os.path.basename(p) or p
        label = f"[attachment {i}: {name}]"

        if not os.path.exists(p):
            reason = f"file not found at {p}"
            print(f"{label} ERROR: {reason}")
            parts.append(f"{label} ERROR: {reason}")
            continue
        if not os.path.isfile(p):
            reason = "not a regular file (is it a folder?)"
            print(f"{label} ERROR: {reason}")
            parts.append(f"{label} ERROR: {reason}")
            continue
        if not os.access(p, os.R_OK):
            reason = "permission denied"
            print(f"{label} ERROR: {reason}")
            parts.append(f"{label} ERROR: {reason}")
            continue

        ext = os.path.splitext(p)[1].lower()
        try:
            if ext in IMAGE_EXTS:
                vision = _describe_image_via_agent(p)
                if vision and not vision.startswith("ERROR:"):
                    body = f"[visual description]\n{vision}"
                else:
                    # Vision unavailable (no key / HTTP error) — fall back to OCR.
                    print(f"{label} vision unavailable ({vision[:80]}); using OCR")
                    body = _read_image(p)
            else:
                body = read_file(p)
        except PermissionError as e:
            reason = f"permission denied ({e})"
            print(f"{label} ERROR: {reason}")
            parts.append(f"{label} ERROR: {reason}")
            continue
        except Exception as e:
            reason = f"read exception: {type(e).__name__}: {e}"
            print(f"{label} ERROR: {reason}")
            parts.append(f"{label} ERROR: {reason}")
            continue

        parts.append(f"{label}\n{body}")

    return "\n\n".join(parts)