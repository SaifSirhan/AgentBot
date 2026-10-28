import importlib
import importlib.util
import json
import os
import sys
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest import mock

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent.parent))


class SafeDocxTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec("office_editor"),
                             "The supplied office editor is not integrated")
        self.editor = importlib.import_module("office_editor")
        directory = tempfile.TemporaryDirectory(prefix="safe_docx_")
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name) / "report with spaces.docx"
        doc = Document()
        doc.add_paragraph("Draft report")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Status"
        doc.save(self.path)
        self.original = self.path.read_bytes()

    def edit(self, action, data):
        return json.loads(self.editor.safe_edit_docx(self.path, action, data))

    def assert_rejected(self, result, code):
        self.assertIs(result["ok"], False)
        self.assertIs(result["retryable"], False)
        self.assertEqual(result["error"]["code"], code)
        self.assertEqual(self.path.read_bytes(), self.original)
        self.assertEqual(list(self.path.parent.glob(".office-editor-*")), [])

    def test_append_paragraph_preserves_existing_content_and_style(self):
        result = self.edit("append_paragraph", {"text": "Summary", "style": "Heading 1"})
        self.assertIs(result["ok"], True)
        self.assertEqual(result["changes"], 1)
        doc = Document(self.path)
        self.assertEqual([p.text for p in doc.paragraphs], ["Draft report", "Summary"])
        self.assertEqual(doc.paragraphs[-1].style.name, "Heading 1")
        self.assertEqual(doc.tables[0].cell(0, 1).text, "Status")

    def test_replace_across_runs_preserves_formatting(self):
        doc = Document(self.path)
        paragraph = doc.paragraphs[0]
        paragraph.clear()
        paragraph.add_run("Dr").bold = True
        paragraph.add_run("aft report").italic = True
        doc.save(self.path)
        result = self.edit("replace_text", {"old": "Draft", "new": "Final"})
        self.assertIs(result["ok"], True)
        self.assertEqual(result["changes"], 1)
        paragraph = Document(self.path).paragraphs[0]
        self.assertEqual(paragraph.text, "Final report")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertTrue(paragraph.runs[1].italic)

    def test_replace_visits_nested_tables_without_duplicate_merged_cells(self):
        doc = Document(self.path)
        cell = doc.tables[0].cell(0, 0).merge(doc.tables[0].cell(0, 1))
        cell.text = "Draft"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "Draft"
        doc.save(self.path)
        result = self.edit("replace_text", {"old": "Draft", "new": "Final"})
        self.assertIs(result["ok"], True)
        self.assertEqual(result["changes"], 3)
        doc = Document(self.path)
        self.assertEqual(doc.paragraphs[0].text, "Final report")
        self.assertEqual(doc.tables[0].cell(0, 0).text.strip(), "Final")
        self.assertEqual(doc.tables[0].cell(0, 0).tables[0].cell(0, 0).text, "Final")

    def test_add_table_row_uses_zero_based_index(self):
        result = self.edit("add_table_row", {"table_index": 0, "values": ["Alice", "Approved"]})
        self.assertIs(result["ok"], True)
        table = Document(self.path).tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual([cell.text for cell in table.rows[1].cells], ["Alice", "Approved"])

    def test_invalid_action_and_payloads_do_not_change_file(self):
        cases = [
            ("delete", {}, "INVALID_ACTION"),
            ("append_paragraph", [], "INVALID_DATA"),
            ("append_paragraph", {}, "INVALID_DATA"),
            ("append_paragraph", {"text": "x", "extra": True}, "INVALID_DATA"),
            ("append_paragraph", {"text": "\x00"}, "INVALID_DATA"),
            ("append_paragraph", {"text": "x", "style": "No such style"}, "INVALID_DATA"),
            ("replace_text", {"old": "", "new": "x"}, "INVALID_DATA"),
            ("add_table_row", {"table_index": True, "values": ["a", "b"]}, "INVALID_DATA"),
            ("add_table_row", {"table_index": 1, "values": ["a", "b"]}, "INVALID_DATA"),
            ("add_table_row", {"table_index": 0, "values": ["a"]}, "INVALID_DATA"),
        ]
        for action, data, code in cases:
            with self.subTest(action=action, data=data):
                self.assert_rejected(self.edit(action, data), code)

    def test_no_match_leaves_original_bytes_untouched(self):
        self.assert_rejected(self.edit("replace_text", {"old": "Missing", "new": "Final"}),
                             "TEXT_NOT_FOUND")

    def test_identical_replacement_does_not_resave(self):
        result = self.edit("replace_text", {"old": "Draft", "new": "Draft"})
        self.assertIs(result["ok"], True)
        self.assertEqual(result["changes"], 0)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_clearing_break_is_rejected_without_losing_layout(self):
        doc = Document(self.path)
        run = doc.paragraphs[0].runs[0]
        run.add_break()
        run._r.find(qn("w:br")).set(qn("w:clear"), "all")
        doc.save(self.path)
        self.original = self.path.read_bytes()
        self.assert_rejected(self.edit("replace_text", {"old": "Draft", "new": "Final"}),
                             "UNSUPPORTED_CONTENT")

    def test_concurrent_external_save_is_preserved(self):
        save = DocumentClass.save
        external_bytes = []

        def save_with_external_edit(document, output):
            save(document, output)
            other = Document(self.path)
            other.add_paragraph("External edit")
            save(other, self.path)
            external_bytes.append(self.path.read_bytes())

        with mock.patch.object(DocumentClass, "save", save_with_external_edit):
            result = self.edit("append_paragraph", {"text": "Our edit"})
        self.assertIs(result["ok"], False)
        self.assertIs(result["retryable"], False)
        self.assertEqual(result["error"]["code"], "CONCURRENT_MODIFICATION")
        self.assertEqual(self.path.read_bytes(), external_bytes[0])
        self.assertEqual([p.text for p in Document(self.path).paragraphs],
                         ["Draft report", "External edit"])
        self.assertEqual(list(self.path.parent.glob(".office-editor-*")), [])

    def test_same_process_writers_do_not_lose_edits(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda n: self.edit("append_paragraph", {"text": f"Edit {n}"}),
                                    range(4)))
        self.assertTrue(all(result["ok"] is True for result in results))
        self.assertCountEqual([p.text for p in Document(self.path).paragraphs],
                              ["Draft report", "Edit 0", "Edit 1", "Edit 2", "Edit 3"])

    def test_failed_atomic_replace_leaves_original_and_removes_temporary_file(self):
        with mock.patch.object(self.editor.os, "replace", side_effect=PermissionError):
            result = self.edit("append_paragraph", {"text": "Not saved"})
        self.assert_rejected(result, "FILE_LOCKED_OR_ACCESS_DENIED")

    @unittest.skipUnless(os.name == "nt", "Windows file-sharing semantics")
    def test_real_windows_lock_prevents_overwrite(self):
        import ctypes
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        create = kernel32.CreateFileW
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create.restype = wintypes.HANDLE
        close = kernel32.CloseHandle
        close.argtypes = [wintypes.HANDLE]
        close.restype = wintypes.BOOL
        handle = create(str(self.path), 0x80000000, 1, None, 3, 0x80, None)
        self.assertNotEqual(handle, wintypes.HANDLE(-1).value)
        try:
            result = self.edit("append_paragraph", {"text": "Locked edit"})
            self.assert_rejected(result, "FILE_LOCKED_OR_ACCESS_DENIED")
        finally:
            self.assertTrue(close(handle))

    def test_missing_and_corrupt_documents_return_json_errors(self):
        missing = self.path.parent / "missing.docx"
        result = json.loads(self.editor.safe_edit_docx(missing, "append_paragraph", {"text": "x"}))
        self.assertEqual(result["error"]["code"], "FILE_NOT_FOUND")
        self.assertFalse(missing.exists())
        broken = self.path.parent / "broken.docx"
        broken.write_bytes(b"not a document")
        result = json.loads(self.editor.safe_edit_docx(broken, "append_paragraph", {"text": "x"}))
        self.assertEqual(result["error"]["code"], "INVALID_DOCUMENT")
        self.assertEqual(broken.read_bytes(), b"not a document")


class SafeDocxDispatchTests(unittest.TestCase):
    def setUp(self):
        import agent
        self.agent = agent
        directory = tempfile.TemporaryDirectory(prefix="safe_docx_dispatch_")
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name) / "report.docx"
        doc = Document()
        doc.add_paragraph("Original")
        doc.save(self.path)
        self.original = self.path.read_bytes()
        group_mode = mock.patch.object(agent, "GROUP_MODE", False)
        group_mode.start()
        self.addCleanup(group_mode.stop)

    def test_parsed_object_input_reaches_safe_editor(self):
        action, error = self.agent.parse_action(json.dumps({
            "tool": "edit_docx_safe", "input": {
                "filepath": str(self.path), "action": "append_paragraph", "data": {"text": "Appended"},
            },
        }))
        self.assertIsNone(error)
        result = json.loads(self.agent.execute_tool(action))
        self.assertIs(result["ok"], True)
        self.assertEqual(result["filepath"], str(self.path))
        self.assertEqual([p.text for p in Document(self.path).paragraphs], ["Original", "Appended"])

    def test_json_string_input_reaches_safe_editor(self):
        result = self.agent.execute_tool({"tool": "edit_docx_safe", "input": json.dumps({
            "filepath": str(self.path), "action": "replace_text", "data": {"old": "Original", "new": "Final"},
        })})
        self.assertFalse(result.startswith("ERROR:"), result)
        self.assertIs(json.loads(result)["ok"], True)
        self.assertEqual(Document(self.path).paragraphs[0].text, "Final")

    def test_invalid_envelopes_return_nonretryable_json_errors(self):
        for value in ("not json", "null", "[]", {}, [], None, 123,
                      {"filepath": str(self.path), "action": "append_paragraph"},
                      {"filepath": str(self.path), "action": "append_paragraph", "data": {}, "extra": 1}):
            with self.subTest(value=value):
                result = self.agent.execute_tool({"tool": "edit_docx_safe", "input": value})
                self.assertFalse(result.startswith("ERROR:"), result)
                parsed = json.loads(result)
                self.assertIs(parsed["ok"], False)
                self.assertIs(parsed["retryable"], False)
                self.assertEqual(parsed["error"]["code"], "INVALID_DATA")
                self.assertEqual(self.path.read_bytes(), self.original)

    def test_module_error_is_forwarded_without_claiming_success(self):
        result = self.agent.execute_tool({"tool": "edit_docx_safe", "input": {
            "filepath": str(self.path), "action": "replace_text", "data": {"old": "Absent", "new": "Final"},
        }})
        self.assertFalse(result.startswith("ERROR:"), result)
        self.assertEqual(json.loads(result)["error"]["code"], "TEXT_NOT_FOUND")
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_group_chat_cannot_edit_document(self):
        with mock.patch.object(self.agent, "GROUP_MODE", True):
            result = self.agent.execute_tool({"tool": "edit_docx_safe", "input": {
                "filepath": str(self.path), "action": "append_paragraph", "data": {"text": "Blocked"},
            }})
        self.assertIn("disabled in group chats", result)
        self.assertEqual(self.path.read_bytes(), self.original)


if __name__ == "__main__":
    unittest.main()
