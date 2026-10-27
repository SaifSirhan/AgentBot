"""Tests for file_tools.inspect_docx, using real python-docx fixtures.

Run: python -m unittest tests.test_inspect_docx -v
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import file_tools
from docx import Document


class InspectDocxTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory(prefix="inspect_docx_")
        self.addCleanup(self.dir.cleanup)

    def path(self, name):
        return os.path.join(self.dir.name, name)

    def build(self, name="sample.docx", paragraphs=(), table=None):
        doc = Document()
        for text, style in paragraphs:
            doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
        if table:
            rows, cols = table
            t = doc.add_table(rows=rows, cols=cols)
            for r in range(rows):
                for c in range(cols):
                    t.cell(r, c).text = f"r{r}c{c}"
        path = self.path(name)
        doc.save(path)
        return path

    def test_structure_is_valid_json(self):
        path = self.build(
            paragraphs=[("Report title", "Heading 1"), ("Body text", None)],
            table=(2, 3),
        )
        payload = json.loads(file_tools.inspect_docx(path))

        self.assertEqual(payload["paragraph_count"], 2)
        self.assertEqual(payload["shown_paragraphs"], 2)
        self.assertEqual(payload["paragraphs"][0]["text"], "Report title")
        self.assertEqual(payload["paragraphs"][0]["style"], "Heading 1")
        self.assertEqual(payload["table_count"], 1)
        table = payload["tables"][0]
        self.assertEqual(table["rows"], 2)
        self.assertEqual(table["columns"], 3)
        self.assertEqual(table["cells"][0][0], "r0c0")
        self.assertEqual(table["cells"][1][2], "r1c2")
        self.assertFalse(payload["truncated"])

    def test_missing_file(self):
        out = file_tools.inspect_docx(self.path("nope.docx"))
        self.assertTrue(out.startswith("ERROR: file not found"))

    def test_rejects_non_docx(self):
        path = self.path("notes.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("hello")
        self.assertIn("only reads .docx", file_tools.inspect_docx(path))

    def test_corrupt_docx_is_reported(self):
        path = self.path("broken.docx")
        with open(path, "wb") as f:
            f.write(b"not a zip")
        self.assertTrue(file_tools.inspect_docx(path).startswith("ERROR: cannot open"))

    def test_empty_document(self):
        path = self.build(paragraphs=[])
        payload = json.loads(file_tools.inspect_docx(path))
        self.assertEqual(payload["paragraph_count"], 0)
        self.assertEqual(payload["tables"], [])

    def test_paragraph_cap_and_count(self):
        path = self.build(paragraphs=[(f"line {i}", None) for i in range(300)])
        payload = json.loads(file_tools.inspect_docx(path))
        self.assertEqual(payload["paragraph_count"], 300)
        self.assertEqual(payload["shown_paragraphs"], file_tools.MAX_INSPECT_PARAGRAPHS)
        self.assertEqual(len(payload["paragraphs"]), file_tools.MAX_INSPECT_PARAGRAPHS)

    def test_long_paragraph_text_is_clipped(self):
        path = self.build(paragraphs=[("y" * 600, None)])
        payload = json.loads(file_tools.inspect_docx(path))
        text = payload["paragraphs"][0]["text"]
        self.assertEqual(len(text), file_tools.MAX_PARAGRAPH_CHARS + 1)
        self.assertTrue(text.endswith("\u2026"))


if __name__ == "__main__":
    unittest.main()
