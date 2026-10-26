"""
rag_tool.py — Local document search for AgentBot.

ChromaDB stores embeddings for semantic search; a SQLite FTS5 table stores the
same chunks for keyword search. Results from both are merged with Reciprocal
Rank Fusion so exact names, IDs and filenames are not lost to vector search.
"""

import os
import json
import glob
import re
import sqlite3
import threading

import chromadb
from sentence_transformers import SentenceTransformer
from pypdf import PdfReader

# Where the vector database lives
DB_PATH = os.path.join(
    os.environ.get("APPDATA", "."),
    "AgentBot",
    "rag_db"
)

# Which embedding model to use. Small and fast.
MODEL_NAME = "all-MiniLM-L6-v2"

STATE_FILE = os.path.join(
    os.environ.get("APPDATA", "."), "AgentBot", "rag_state.json"
)

# Root of the flattened Telegram export. Kept in sync with
# telegram_export_parser.OUTPUT_DIR; consumers use it to scope searches to the
# group chat rather than every indexed folder.
GROUP_EXPORT_ROOT = r"C:\Users\USER\PrivateExport"

# Reciprocal Rank Fusion constant. A chunk that appears in both the semantic
# and keyword lists receives both 1/(k+rank) terms, which puts it far above
# single-list hits.
RRF_K = 60


def _filename_for(source: str) -> str:
    return source.replace("\\", "/").rsplit("/", 1)[-1]


def rrf_scores(rankings, k=RRF_K):
    """Return {chunk_id: score} for one or more ranked id lists.

    Each list contributes 1/(k + rank) to every id it contains. Overlap
    between lists is therefore additive: a chunk found by both Chroma and
    FTS5 outranks a chunk found by only one of them.
    """
    scores = {}
    for ranking in rankings:
        for rank, chunk_id in enumerate(ranking, start=1):
            scores[chunk_id] = scores.get(chunk_id, 0.0) + 1.0 / (k + rank)
    return scores


class FTSIndex:
    """SQLite FTS5 keyword index for RAG chunks.

    The filename is indexed separately from the body so a query like
    ``Q4_report_2026.pdf`` finds the file even when its text never mentions
    the name. One connection is shared by GUI and agent threads, so every
    operation takes a lock.
    """

    def __init__(self, path: str):
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        self._lock = threading.Lock()
        self._conn = sqlite3.connect(path, check_same_thread=False)
        self._conn.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(
                id UNINDEXED,
                source UNINDEXED,
                filename,
                text,
                tokenize = 'unicode61'
            )
            """
        )
        self._conn.commit()

    def close(self):
        with self._lock:
            try:
                self._conn.close()
            except Exception:
                pass

    def count(self) -> int:
        with self._lock:
            row = self._conn.execute(
                "SELECT COUNT(*) FROM chunks").fetchone()
        return int(row[0]) if row else 0

    def replace_source(self, source: str, rows):
        """Replace every chunk for ``source`` with ``(chunk_id, text)`` rows."""
        filename = _filename_for(source)
        with self._lock:
            self._conn.execute(
                "DELETE FROM chunks WHERE source = ?", (source,))
            self._conn.executemany(
                """
                INSERT INTO chunks (id, source, filename, text)
                VALUES (?, ?, ?, ?)
                """,
                [(chunk_id, source, filename, text)
                 for chunk_id, text in rows],
            )
            self._conn.commit()

    def bulk_insert(self, rows):
        """Insert ``(chunk_id, source, text)`` rows without deduplicating."""
        with self._lock:
            self._conn.executemany(
                """
                INSERT INTO chunks (id, source, filename, text)
                VALUES (?, ?, ?, ?)
                """,
                [(chunk_id, source, _filename_for(source), text)
                 for chunk_id, source, text in rows],
            )
            self._conn.commit()

    def delete_source(self, source: str):
        with self._lock:
            self._conn.execute(
                "DELETE FROM chunks WHERE source = ?", (source,))
            self._conn.commit()

    def search(self, query: str, limit: int = 10):
        """Return ranked keyword hits as dicts with id/source/text.

        Query text is tokenised to word characters and quoted, so FTS5
        operators in user input are searched as plain words instead of
        changing the query grammar.
        """
        terms = re.findall(r"\w+", query or "", flags=re.UNICODE)
        if not terms:
            return []
        match = " OR ".join(f'"{term}"' for term in terms)
        limit = max(1, int(limit))
        with self._lock:
            try:
                rows = self._conn.execute(
                    """
                    SELECT id, source, text
                    FROM chunks
                    WHERE chunks MATCH ?
                    ORDER BY rank
                    LIMIT ?
                    """,
                    (match, limit),
                ).fetchall()
            except sqlite3.Error as e:
                print(f"[rag] FTS query failed: {e}")
                return []
        return [
            {"id": chunk_id, "source": source, "text": text}
            for chunk_id, source, text in rows
        ]


def _load_state() -> dict:
    import json
    if not os.path.exists(STATE_FILE):
        return {}
    try:
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _save_state(state: dict):
    import json
    try:
        os.makedirs(os.path.dirname(STATE_FILE), exist_ok=True)
        with open(STATE_FILE, "w", encoding="utf-8") as f:
            json.dump(state, f, indent=2)
    except Exception as e:
        print(f"[rag] could not save state: {e}")


def _file_signature(path: str):
    """Return (mtime, size) or None if file is gone."""
    try:
        st = os.stat(path)
        return (st.st_mtime, st.st_size)
    except OSError:
        return None
    
class RAGTool:
    def __init__(self, db_path: str = None, fts_path: str = None):
        db_path = db_path or DB_PATH
        os.makedirs(db_path, exist_ok=True)
        self.client = chromadb.PersistentClient(path=db_path)
        self.collection = self.client.get_or_create_collection(
            name="documents",
            metadata={"hnsw:space": "cosine"}
        )
        self.model = SentenceTransformer(MODEL_NAME)
        self.fts = FTSIndex(
            fts_path or os.path.join(db_path, "rag_fts.sqlite3"))
        self._backfill_fts()

    def _backfill_fts(self):
        """Migrate an existing Chroma-only index into FTS once."""
        if self.fts is None or self.fts.count() != 0:
            return
        try:
            total = self.collection.count()
        except Exception:
            return
        if total == 0:
            return
        try:
            data = self.collection.get(include=["documents", "metadatas"])
        except Exception as e:
            print(f"[rag] FTS backfill skipped: {e}")
            return
        ids = data.get("ids") or []
        documents = data.get("documents") or []
        metadatas = data.get("metadatas") or []
        rows = []
        for chunk_id, text, meta in zip(ids, documents, metadatas):
            if not text:
                continue
            source = (meta or {}).get("source", "unknown")
            rows.append((chunk_id, source, text))
        if not rows:
            return
        try:
            self.fts.bulk_insert(rows)
            print(f"[rag] backfilled {len(rows)} chunks into FTS5")
        except Exception as e:
            print(f"[rag] FTS backfill failed: {e}")

    # ----------------------------------------------------------
    # INDEXING
    # ----------------------------------------------------------
    def index_file(self, filepath: str, chunk_size: int = None,
                   overlap: int = None) -> str:
        """Read one file, chunk it, embed it, store in ChromaDB.

        chunk_size/overlap default to None so callers can pass explicit
        values; when omitted they are chosen from the filepath.
        """
        ext = os.path.splitext(filepath)[1].lower()

        if ext == ".pdf":
            text = self._read_pdf(filepath)
            if not text or len(text.strip()) < 50:
                return f"Skipped {filepath} (scanned PDF or no text layer)"
        elif ext == ".docx":
            try:
                text = self._read_docx(filepath)
            except Exception as e:
                return f"Skipped {filepath} (docx read failed: {e})"
        elif ext == ".xlsx":
            try:
                text = self._read_xlsx(filepath)
            except Exception as e:
                return f"Skipped {filepath} (xlsx read failed: {e})"
        elif ext in (
            ".txt", ".md", ".py", ".json", ".log",
            ".php", ".js", ".ts", ".jsx", ".tsx",
            ".html", ".css", ".scss",
            ".yaml", ".yml", ".toml", ".ini", ".cfg",
            ".sh", ".bat", ".ps1", ".sql", ".java", ".c", ".cpp", ".go", ".rs",
        ):
            try:
                with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            except Exception as e:
                return f"Skipped {filepath} (read failed: {e})"
        else:
            return f"Skipped {filepath} (unsupported type)"

        if not text or not text.strip():
            return f"Skipped {filepath} (empty)"

        # Skip huge files
        if len(text) > 500_000:
            return f"Skipped {filepath} (too large: {len(text)} chars)"

        # Chat logs need fine-grained chunks: a 1000-char window spans 8-15
        # messages, so a chunk's date/author blurs across several messages.
        use_small = "PrivateExport" in filepath
        if chunk_size is None:
            chunk_size = 300 if use_small else 1000
        if overlap is None:
            overlap = 50 if use_small else 200

        chunks = self._chunk(text, size=chunk_size, overlap=overlap)
        if not chunks:
            return f"No text chunks from {filepath}"

        ids = [f"{filepath}::{i}" for i in range(len(chunks))]
        embeddings = self.model.encode(chunks, show_progress_bar=False).tolist()

        try:
            self.collection.delete(where={"source": filepath})
        except Exception:
            pass

        self.collection.upsert(
            ids=ids,
            documents=chunks,
            embeddings=embeddings,
            metadatas=[{"source": filepath} for _ in chunks]
        )
        self.fts.replace_source(filepath, zip(ids, chunks))
        return f"Indexed {len(chunks)} chunks from {filepath}"

    def index_folder(self, folder: str, chunk_size: int = None,
                     overlap: int = None) -> str:
        """Index new/changed files only. Delete chunks for removed files."""
        SUPPORTED = {
            ".pdf", ".txt", ".md", ".py", ".json", ".log",
            ".php", ".js", ".ts", ".jsx", ".tsx",
            ".html", ".css", ".scss",
            ".yaml", ".yml", ".toml", ".ini", ".cfg",
            ".sh", ".bat", ".ps1", ".sql", ".java", ".c", ".cpp", ".go", ".rs",
            ".docx", ".xlsx",
        }

        SKIP_DIRS = {
            ".git", "node_modules", "__pycache__", ".venv", "venv",
            "env", "dist", "build", ".next", ".cache", "vendor",
        }

        SKIP_NAMES = {
            "api key.txt", "password.txt", "secrets.txt",
            "github-recovery-codes.txt",
        }

        SKIP_TOKENS = (
            "api_key", "apikey", "secret", "password",
            "token", "recovery", "credential",
        )

        state = _load_state()
        seen_paths = set()
        indexed, skipped, removed, failed = [], [], [], []

        # ---- Pass 1: walk and index changed/new files ----
        for root, dirs, files in os.walk(folder):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS]

            for name in files:
                ext = os.path.splitext(name)[1].lower()
                if ext not in SUPPORTED:
                    continue
                if name.lower() in SKIP_NAMES:
                    continue
                if any(tok in name.lower() for tok in SKIP_TOKENS):
                    continue

                path = os.path.join(root, name)
                seen_paths.add(path)

                sig = _file_signature(path)
                if sig is None:
                    continue

                prev = state.get(path)
                if prev and prev.get("mtime") == sig[0] and prev.get("size") == sig[1]:
                    # File unchanged since last index — skip
                    skipped.append(path)
                    continue

                print(f"  [new/changed] {path}", flush=True)
                try:
                    result = self.index_file(path, chunk_size=chunk_size,
                                             overlap=overlap)
                    if result.startswith("Indexed"):
                        state[path] = {"mtime": sig[0], "size": sig[1]}
                        indexed.append(path)
                    else:
                        failed.append(f"{path}: {result}")
                except Exception as e:
                    failed.append(f"{path}: {e}")

        # ---- Pass 2: remove DB entries for files that no longer exist ----
        for path in list(state.keys()):
            if path in seen_paths:
                continue
            # Only touch paths that were inside the folder we just walked
            if not os.path.abspath(path).startswith(os.path.abspath(folder)):
                continue
            try:
                self.collection.delete(where={"source": path})
            except Exception:
                pass
            try:
                self.fts.delete_source(path)
            except Exception:
                pass
            removed.append(path)
            state.pop(path, None)

        _save_state(state)

        # ---- Build summary ----
        summary = (
            f"Indexed: {len(indexed)} new/changed\n"
            f"Skipped (unchanged): {len(skipped)}\n"
            f"Removed (deleted): {len(removed)}\n"
            f"Failed: {len(failed)}"
        )
        if indexed:
            summary += "\n\nNew/changed:\n" + "\n".join(
                f"  + {os.path.basename(p)}" for p in indexed[:10]
            )
            if len(indexed) > 10:
                summary += f"\n  ...and {len(indexed) - 10} more"
        if removed:
            summary += "\n\nRemoved:\n" + "\n".join(
                f"  - {os.path.basename(p)}" for p in removed[:10]
            )
        if failed:
            summary += "\n\nFailed:\n" + "\n".join(f"  ! {f}" for f in failed[:5])

        return summary

    # ----------------------------------------------------------
    # SEARCH
    # ----------------------------------------------------------
    def search(self, query: str, top_k: int = 8) -> str:
        """Find relevant chunks using semantic + keyword search, fused by RRF."""
        top_k = max(1, int(top_k))
        candidate_k = top_k * 2
        semantic = self._semantic_results(query, candidate_k)
        keyword = self._keyword_results(query, candidate_k)

        if not semantic and not keyword:
            text_count = 0
            fts_count = 0
            try:
                text_count = self.collection.count()
            except Exception:
                pass
            if self.fts is not None:
                try:
                    fts_count = self.fts.count()
                except Exception:
                    pass
            if text_count == 0 and fts_count == 0:
                return "No documents indexed yet."
            return "No relevant information found."

        documents = {}
        for hit in semantic + keyword:
            documents.setdefault(hit["id"], hit)

        scores = rrf_scores([
            [hit["id"] for hit in semantic],
            [hit["id"] for hit in keyword],
        ])
        ordered = sorted(
            scores, key=lambda chunk_id: (-scores[chunk_id], chunk_id))

        output_parts = []
        for chunk_id in ordered[:top_k]:
            hit = documents[chunk_id]
            output_parts.append(
                f"[From: {hit.get('source', 'unknown')}]\n{hit['text']}")
        return "\n\n---\n\n".join(output_parts)

    def _semantic_results(self, query: str, limit: int):
        if not query or not query.strip():
            return []
        try:
            if self.collection.count() == 0:
                return []
        except Exception:
            return []
        query_embedding = self.model.encode(
            [query], show_progress_bar=False).tolist()
        results = self.collection.query(
            query_embeddings=query_embedding,
            n_results=limit,
        )
        ids = (results.get("ids") or [[]])[0] or []
        documents = (results.get("documents") or [[]])[0] or []
        metadatas = (results.get("metadatas") or [[]])[0] or []
        hits = []
        for chunk_id, text, meta in zip(ids, documents, metadatas):
            if not text:
                continue
            hits.append({
                "id": chunk_id,
                "source": (meta or {}).get("source", "unknown"),
                "text": text,
            })
        return hits

    def _keyword_results(self, query: str, limit: int):
        if self.fts is None or not query or not query.strip():
            return []
        try:
            return self.fts.search(query, limit=limit)
        except Exception as e:
            print(f"[rag] keyword search failed: {e}")
            return []

    # ----------------------------------------------------------
    # HELPERS
    # ----------------------------------------------------------
    def _read_pdf(self, path: str) -> str:
        reader = PdfReader(path)
        parts = []
        for page in reader.pages:
            try:
                t = page.extract_text() or ""
                parts.append(t)
            except Exception:
                continue
        return "\n".join(parts)

    def _read_docx(self, path: str) -> str:
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    def _read_xlsx(self, path: str) -> str:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)

    def _chunk(self, text: str, size: int = 1000, overlap: int = 200):
        """Split text into overlapping chunks."""
        chunks = []
        start = 0
        while start < len(text):
            end = start + size
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            start = end - overlap
        return chunks


# ----------------------------------------------------------
# Singleton
# ----------------------------------------------------------
_rag = None

def get_rag():
    global _rag
    if _rag is None:
        _rag = RAGTool()
    return _rag


def index_documents(path: str) -> str:
    """Public function AgentBot calls."""
    rag = get_rag()
    if os.path.isdir(path):
        return rag.index_folder(path)
    return rag.index_file(path)


def search_documents(query: str) -> str:
    """Public function AgentBot calls."""
    return get_rag().search(query)