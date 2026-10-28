import json
import os
import stat
import tempfile
import threading
from pathlib import Path

__all__ = ["safe_edit_docx"]
_EDIT_LOCK = threading.Lock()


def safe_edit_docx(filepath, action, data):
    """Edit an existing DOCX in place; return a JSON string without retrying."""
    temporary_path = None

    def error(code, message):
        return json.dumps({
            "ok": False,
            "retryable": False,
            "error": {"code": code, "message": message},
        })

    def require_text(value, name):
        if not isinstance(value, str):
            raise ValueError(f"{name} must be a string.")
        if any(not (
            char in "\t\n\r"
            or 0x20 <= ord(char) <= 0xD7FF
            or 0xE000 <= ord(char) <= 0xFFFD
            or 0x10000 <= ord(char) <= 0x10FFFF
        ) for char in value):
            raise ValueError(f"{name} contains characters forbidden in DOCX text.")

    def paragraphs(container, seen):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        yield from paragraphs(cell, seen)

    def replace_in_runs(paragraph, old, new):
        spans = []
        offset = 0
        for run in paragraph.runs:
            end = offset + len(run.text)
            if end > offset:
                spans.append((offset, end, run))
            offset = end

        text = paragraph.text
        matches = []
        position = text.find(old)
        while position != -1:
            matches.append(position)
            position = text.find(old, position + len(old))

        # Right-to-left edits keep the original offsets valid for earlier matches.
        for start in reversed(matches):
            end = start + len(old)
            touched = [span for span in spans if span[0] < end and span[1] > start]
            first_start, _, first = touched[0]
            last_start, _, last = touched[-1]
            prefix = first.text[:start - first_start]
            suffix = last.text[end - last_start:]
            if first is last:
                first.text = prefix + new + suffix
            else:
                first.text = prefix + new
                for _, _, run in touched[1:-1]:
                    run.text = ""
                last.text = suffix
        return len(matches)

    _EDIT_LOCK.acquire()
    try:
        try:
            from docx import Document
            from docx.enum.style import WD_STYLE_TYPE
            from docx.oxml.ns import qn
        except ImportError:
            return error("MISSING_DEPENDENCY", "Install python-docx: python -m pip install python-docx")

        schemas = {
            "append_paragraph": ({"text"}, {"style"}),
            "replace_text": ({"old", "new"}, set()),
            "add_table_row": ({"table_index", "values"}, set()),
        }
        if not isinstance(action, str) or action not in schemas:
            return error("INVALID_ACTION", "Use append_paragraph, replace_text, or add_table_row.")
        if not isinstance(data, dict):
            return error("INVALID_DATA", "data must be a dictionary.")
        required, optional = schemas[action]
        if not required <= data.keys() or not data.keys() <= required | optional:
            return error("INVALID_DATA", f"Required keys: {sorted(required)}; optional keys: {sorted(optional)}.")

        if action == "append_paragraph":
            require_text(data["text"], "text")
            if "style" in data:
                require_text(data["style"], "style")
                if not data["style"]:
                    raise ValueError("style must not be empty.")
        elif action == "replace_text":
            require_text(data["old"], "old")
            require_text(data["new"], "new")
            if not data["old"]:
                raise ValueError("old must not be empty.")
        else:
            if type(data["table_index"]) is not int or data["table_index"] < 0:
                raise ValueError("table_index must be a nonnegative integer (zero-based).")
            if not isinstance(data["values"], list):
                raise ValueError("values must be a list of strings.")
            for value in data["values"]:
                require_text(value, "Each table value")

        try:
            path = Path(filepath).absolute()
            if "\x00" in str(path) or path.suffix.lower() != ".docx":
                return error("INVALID_PATH", "filepath must identify an existing .docx file.")
            if path.is_symlink():
                return error("INVALID_PATH", "Use the document's real path, not a symbolic link.")
        except (TypeError, ValueError):
            return error("INVALID_PATH", "filepath must be a .docx path string or Path object.")

        with path.open("rb") as source:
            original_stat = os.fstat(source.fileno())
            try:
                document = Document(source)
            except OSError:
                raise
            except Exception:
                return error("INVALID_DOCUMENT", "The file is not a readable, unencrypted Word DOCX document.")

        if action == "append_paragraph":
            style = None
            if "style" in data:
                try:
                    style = document.styles[data["style"]]
                except KeyError:
                    raise ValueError("The requested paragraph style does not exist.") from None
                if style.type != WD_STYLE_TYPE.PARAGRAPH:
                    raise ValueError("style must name a paragraph style.")
            document.add_paragraph(data["text"], style=style)
            changes = 1

        elif action == "replace_text":
            old, new = data["old"], data["new"]
            changes = 0
            matched = False
            paragraph_tags = {qn(f"w:{name}") for name in (
                "pPr", "r", "bookmarkStart", "bookmarkEnd", "proofErr",
            )}
            run_tags = {qn(f"w:{name}") for name in ("rPr", "t", "tab", "br", "cr")}
            for paragraph in paragraphs(document, set()):
                if old not in paragraph.text:
                    continue
                matched = True
                if old == new:
                    continue
                # Run.text destroys non-text children; inspect only, never rewrite XML.
                if (
                    any(child.tag not in paragraph_tags for child in paragraph._p)
                    or any(child.tag not in run_tags for run in paragraph.runs for child in run._r)
                    or any(
                        child.tag == qn("w:br") and (
                            child.get(qn("w:type")) not in (None, "textWrapping")
                            or child.get(qn("w:clear")) not in (None, "none")
                        )
                        for run in paragraph.runs for child in run._r
                    )
                    or paragraph.text != "".join(run.text for run in paragraph.runs)
                ):
                    return error("UNSUPPORTED_CONTENT", "A matching paragraph contains non-plain content; no changes were saved.")
                changes += replace_in_runs(paragraph, old, new)
            if not matched:
                return error("TEXT_NOT_FOUND", "No literal match in supported body paragraphs or table cells.")

        else:
            index, values = data["table_index"], data["values"]
            if index >= len(document.tables):
                raise ValueError(f"table_index is out of range; the document has {len(document.tables)} top-level tables.")
            table = document.tables[index]
            if len(values) != len(table.columns):
                raise ValueError(f"values must contain exactly {len(table.columns)} strings.")
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
            changes = 1

        if changes:
            with tempfile.NamedTemporaryFile(
                mode="w+b", prefix=".office-editor-", suffix=".docx",
                dir=path.parent, delete=False,
            ) as output:
                temporary_path = Path(output.name)
                document.save(output)
                output.flush()
                os.fsync(output.fileno())
            if os.name != "nt":
                os.chmod(temporary_path, stat.S_IMODE(original_stat.st_mode))
            current_stat = path.stat()
            if any(
                getattr(current_stat, field) != getattr(original_stat, field)
                for field in ("st_dev", "st_ino", "st_size", "st_mtime_ns", "st_ctime_ns")
            ):
                return error("CONCURRENT_MODIFICATION", "The document changed during editing; no changes from this call were saved.")
            os.replace(temporary_path, path)
            temporary_path = None

        return json.dumps({"ok": True, "action": action, "filepath": str(path), "changes": changes})

    except PermissionError:
        return error("FILE_LOCKED_OR_ACCESS_DENIED", "Close the file in Microsoft Word and check file/folder write permissions before trying again.")
    except FileNotFoundError:
        return error("FILE_NOT_FOUND", "The document or its containing folder does not exist.")
    except IsADirectoryError:
        return error("INVALID_PATH", "filepath refers to a directory, not a document.")
    except ValueError as exc:
        return error("INVALID_DATA", str(exc))
    except OSError as exc:
        if getattr(exc, "winerror", None) in (32, 33):
            return error("FILE_LOCKED_OR_ACCESS_DENIED", "Close the file in Microsoft Word before trying again.")
        return error("IO_ERROR", f"The document could not be read or saved (OS error {exc.errno}).")
    except Exception:
        return error("EDIT_FAILED", "The edit failed; no successful save was reported.")
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass
        _EDIT_LOCK.release()
